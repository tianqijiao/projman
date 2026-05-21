from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain import AttachmentKind
from app.services import ProjectOverview


EXPORT_COLUMNS = [
    "年度",
    "项目名称",
    "预算金额",
    "合同金额",
    "预算差额",
    "合同开始",
    "合同结束",
    "验收日期",
    "付款日期",
    "状态",
    "缺少材料",
    "采购依据",
    "合同审签 PDF",
    "盖章合同扫描件",
    "验收单",
    "发票",
    "其他附件",
    "备注",
]


def build_projects_excel(overviews: list[ProjectOverview], title: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "项目台账"

    sheet.append([title])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(EXPORT_COLUMNS))
    title_cell = sheet.cell(row=1, column=1)
    title_cell.font = Font(size=16, bold=True, color="1F2937")
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet.append(EXPORT_COLUMNS)
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="243B53")
        cell.fill = PatternFill("solid", fgColor="EAF1F8")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in overviews:
        project = item.project
        attachments_by_kind = {
            kind: [
                attachment.original_filename
                for attachment in item.attachments
                if attachment.kind == kind
            ]
            for kind in AttachmentKind
        }
        sheet.append(
            [
                project.year,
                project.name,
                project.budget_amount,
                project.contract_amount,
                item.budget_delta.difference,
                project.contract_start,
                project.contract_end,
                project.acceptance_date,
                project.payment_date,
                _status_text(item),
                "、".join(item.status.missing_evidence),
                _join_files(attachments_by_kind[AttachmentKind.PROCUREMENT_BASIS]),
                _join_files(attachments_by_kind[AttachmentKind.CONTRACT_REVIEW]),
                _join_files(attachments_by_kind[AttachmentKind.SIGNED_CONTRACT]),
                _join_files(attachments_by_kind[AttachmentKind.ACCEPTANCE]),
                _join_files(attachments_by_kind[AttachmentKind.INVOICE]),
                _join_files(attachments_by_kind[AttachmentKind.OTHER]),
                project.notes,
            ]
        )

    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(EXPORT_COLUMNS))}{max(sheet.max_row, 2)}"
    _format_columns(sheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _status_text(item: ProjectOverview) -> str:
    labels = []
    if item.status.established:
        labels.append("已立项")
    if item.status.contract_signed:
        labels.append("合同已签")
    if item.status.accepted:
        labels.append("已验收")
    if item.status.paid:
        labels.append("已付款")
    return "、".join(labels) or "未立项"


def _join_files(files: list[str]) -> str:
    return "；".join(files)


def _format_columns(sheet) -> None:
    widths = {
        "A": 10,
        "B": 26,
        "C": 14,
        "D": 14,
        "E": 14,
        "F": 14,
        "G": 14,
        "H": 14,
        "I": 14,
        "J": 24,
        "K": 24,
        "L": 28,
        "M": 28,
        "N": 28,
        "O": 28,
        "P": 28,
        "Q": 28,
        "R": 30,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row in sheet.iter_rows(min_row=3, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '#,##0.00'

    for row in sheet.iter_rows(min_row=3, min_col=6, max_col=9):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"
