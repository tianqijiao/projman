from datetime import date, timedelta
from pathlib import Path
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlmodel import Session, select

from app.domain import AttachmentKind
from app.main import create_app
from app.models import AnnualAttachment, AnnualExecution, Attachment, Project


def make_client(tmp_path: Path) -> TestClient:
    app = create_app(
        database_url=f"sqlite:///{tmp_path / 'app.db'}",
        data_dir=tmp_path / "data",
        username="admin",
        password="secret",
        secret_key="test-secret",
    )
    return TestClient(app)


def login(client: TestClient) -> None:
    response = client.post(
        "/login",
        data={"username": "admin", "password": "secret"},
        follow_redirects=False,
    )
    assert response.status_code == 303


def test_dashboard_requires_login(tmp_path):
    client = make_client(tmp_path)

    response = client.get("/", follow_redirects=False)

    assert response.status_code == 303
    assert response.headers["location"] == "/login"


def test_business_and_attachment_routes_require_login(tmp_path):
    client = make_client(tmp_path)

    protected_paths = [
        "/projects",
        "/projects/1",
        "/attachments/1/preview",
        "/attachments/1/file",
        "/attachments/1/download",
        "/annual-attachments/1/preview",
        "/annual-attachments/1/file",
        "/annual-attachments/1/download",
    ]

    for path in protected_paths:
        response = client.get(path, follow_redirects=False)
        assert response.status_code == 303, path
        assert response.headers["location"] == "/login", path


def test_login_rejects_bad_password_and_logout_protects_pages(tmp_path):
    client = make_client(tmp_path)

    bad_login_response = client.post(
        "/login",
        data={"username": "admin", "password": "wrong"},
        follow_redirects=False,
    )
    login(client)
    logged_in_response = client.get("/projects")
    logout_response = client.post("/logout", follow_redirects=False)
    after_logout_response = client.get("/projects", follow_redirects=False)

    assert bad_login_response.status_code == 401
    assert "账号或密码不正确" in bad_login_response.text
    assert logged_in_response.status_code == 200
    assert logout_response.status_code == 303
    assert logout_response.headers["location"] == "/login"
    assert after_logout_response.status_code == 303
    assert after_logout_response.headers["location"] == "/login"


def test_dashboard_defaults_to_current_year(tmp_path):
    client = make_client(tmp_path)
    login(client)

    response = client.get("/")

    assert response.status_code == 200
    assert f"{date.today().year} 年度执行看板" in response.text


def test_settings_update_changes_renewal_window_on_dashboard(tmp_path):
    client = make_client(tmp_path)
    login(client)
    today = date.today()
    contract_end = today + timedelta(days=30)
    client.post(
        "/projects",
        data={"year": str(today.year), "name": "续采窗口项目", "budget_amount": "30000"},
    )
    client.post(
        "/projects/1/edit",
        data={
            "year": str(today.year),
            "name": "续采窗口项目",
            "budget_amount": "30000",
            "contract_amount": "30000",
            "contract_start": today.isoformat(),
            "contract_end": contract_end.isoformat(),
            "notes": "",
        },
    )

    client.post("/settings", data={"renewal_lead_days": "10"})
    narrow_response = client.get(f"/?year={today.year}")
    client.post("/settings", data={"renewal_lead_days": "40"})
    wide_response = client.get(f"/?year={today.year}")

    assert "暂无提醒期内到期的合同" in narrow_response.text
    assert "启动下一期采购" in wide_response.text


def test_settings_rejects_non_numeric_renewal_days_with_form_validation(tmp_path):
    client = make_client(tmp_path)
    login(client)

    response = client.post(
        "/settings",
        data={"renewal_lead_days": "not-a-number"},
        follow_redirects=False,
    )

    assert response.status_code == 422


@pytest.mark.parametrize("renewal_days", ["0", "-5"])
def test_settings_rejects_non_positive_renewal_days(tmp_path, renewal_days):
    client = make_client(tmp_path)
    login(client)
    client.post("/settings", data={"renewal_lead_days": "30"})

    response = client.post(
        "/settings",
        data={"renewal_lead_days": renewal_days},
        follow_redirects=False,
    )
    settings_response = client.get("/settings")

    assert response.status_code == 400
    assert "续采提醒提前天数必须大于 0" in response.text
    assert 'value="30"' in settings_response.text


def test_project_list_is_first_navigation_item(tmp_path):
    client = make_client(tmp_path)
    login(client)

    response = client.get("/projects")

    assert response.status_code == 200
    assert response.text.index('href="/projects">项目台账') < response.text.index(
        'href="/">年度看板'
    )


def test_new_project_is_not_in_top_navigation_but_kept_in_toolbar(tmp_path):
    client = make_client(tmp_path)
    login(client)

    response = client.get("/projects")

    assert response.status_code == 200
    header_html = response.text.split("</header>", maxsplit=1)[0]
    assert 'href="/projects/new">新增项目' not in header_html
    assert 'href="/projects/new">新增项目' in response.text


def test_new_project_defaults_to_next_calendar_year(tmp_path):
    client = make_client(tmp_path)
    login(client)

    response = client.get("/projects/new")

    assert response.status_code == 200
    assert f'name="year" type="number" value="{date.today().year + 1}"' in response.text


def test_create_project_rejects_blank_name_and_invalid_budget(tmp_path):
    client = make_client(tmp_path)
    login(client)

    blank_name_response = client.post(
        "/projects",
        data={"year": "2027", "name": "   ", "budget_amount": "100"},
        follow_redirects=False,
    )
    invalid_budget_response = client.post(
        "/projects",
        data={"year": "2027", "name": "非法金额项目", "budget_amount": "abc"},
        follow_redirects=False,
    )
    with Session(client.app.state.engine) as session:
        projects = session.exec(select(Project)).all()

    assert blank_name_response.status_code == 400
    assert "项目名称不能为空" in blank_name_response.text
    assert invalid_budget_response.status_code == 400
    assert "预算金额格式不正确" in invalid_budget_response.text
    assert projects == []


@pytest.mark.parametrize("budget_amount", ["inf", "-inf", "1e309"])
def test_create_project_rejects_non_finite_budget_amount(tmp_path, budget_amount):
    client = make_client(tmp_path)
    login(client)

    response = client.post(
        "/projects",
        data={"year": "2027", "name": "非有限金额项目", "budget_amount": budget_amount},
        follow_redirects=False,
    )
    with Session(client.app.state.engine) as session:
        projects = session.exec(select(Project)).all()

    assert response.status_code == 400
    assert "预算金额格式不正确" in response.text
    assert projects == []


def test_login_create_project_and_show_it_on_dashboard(tmp_path):
    client = make_client(tmp_path)
    login(client)

    response = client.post(
        "/projects",
        data={"year": "2027", "name": "行政电脑维护服务", "budget_amount": "50000"},
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert "行政电脑维护服务" in response.text
    assert "未签合同" in response.text


def test_edit_project_and_upload_pdf_attachment(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "行政电脑维护服务", "budget_amount": "50000"},
    )

    edit_response = client.post(
        "/projects/1/edit",
        data={
            "year": "2027",
            "name": "行政电脑维护服务",
            "budget_amount": "50000",
            "contract_amount": "48000",
            "contract_start": "2027-01-01",
            "contract_end": "2027-12-31",
            "acceptance_date": "",
            "payment_date": "",
            "notes": "测试备注",
        },
        follow_redirects=True,
    )
    upload_response = client.post(
        "/projects/1/attachments",
        data={"kind": "signed_contract"},
        files={"file": ("合同.pdf", b"%PDF-1.7 fake", "application/pdf")},
        follow_redirects=True,
    )

    assert edit_response.status_code == 200
    assert upload_response.status_code == 200
    assert "合同已签" in upload_response.text
    assert "合同.pdf" in upload_response.text


def test_project_list_shows_all_attachment_types_and_download_links(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "行政电脑维护服务", "budget_amount": "50000"},
    )
    client.post(
        "/projects/1/attachments",
        data={"kind": "signed_contract"},
        files={"file": ("盖章合同.pdf", b"%PDF-1.7 fake", "application/pdf")},
    )
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()
    client.post(
        f"/annual-executions/{execution.id}/attachments",
        data={"kind": "invoice"},
        files={"file": ("发票.pdf", b"%PDF-1.7 fake", "application/pdf")},
    )

    response = client.get("/projects")

    assert response.status_code == 200
    assert "采购依据" in response.text
    assert "合同审签 PDF" in response.text
    assert "盖章合同扫描件" in response.text
    assert "验收单" in response.text
    assert "发票" in response.text
    assert "盖章合同.pdf" in response.text
    assert "发票.pdf" in response.text
    assert 'href="/attachments/1/preview"' in response.text
    assert 'href="/attachments/1/download"' in response.text
    assert 'href="/annual-attachments/1/preview"' in response.text
    assert 'href="/annual-attachments/1/download"' in response.text
    assert 'href="/projects/1/attachments/download-all"' in response.text
    assert '<details class="attachment-type" open>' not in response.text


def test_attachment_preview_download_and_download_all(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "行政电脑维护服务", "budget_amount": "50000"},
    )
    client.post(
        "/projects/1/attachments",
        data={"kind": "procurement_basis"},
        files={"file": ("采购依据.pdf", b"%PDF-1.7 fake basis", "application/pdf")},
    )
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()
    client.post(
        f"/annual-executions/{execution.id}/attachments",
        data={"kind": "invoice"},
        files={"file": ("发票.pdf", b"%PDF-1.7 fake invoice", "application/pdf")},
    )

    preview_response = client.get("/attachments/1/preview")
    inline_response = client.get("/attachments/1/file")
    download_response = client.get("/attachments/1/download")
    all_response = client.get("/projects/1/attachments/download-all")

    assert preview_response.status_code == 200
    assert "采购依据.pdf" in preview_response.text
    assert 'src="/attachments/1/file"' in preview_response.text
    assert 'href="/attachments/1/download"' in preview_response.text
    assert inline_response.headers["content-disposition"].startswith("inline;")
    assert download_response.headers["content-disposition"].startswith("attachment;")
    assert all_response.status_code == 200
    assert all_response.headers["content-type"].startswith("application/zip")
    assert "project-1-attachments.zip" in all_response.headers["content-disposition"]

    zip_path = tmp_path / "attachments.zip"
    zip_path.write_bytes(all_response.content)
    with ZipFile(zip_path) as attachments_zip:
        names = attachments_zip.namelist()
        assert any(name.endswith("采购依据.pdf") for name in names)
        assert any(name.endswith("发票.pdf") for name in names)


def test_empty_project_attachment_zip_download_is_valid_empty_zip(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "暂无附件项目", "budget_amount": "50000"},
    )

    response = client.get("/projects/1/attachments/download-all")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/zip")
    assert "project-1-attachments.zip" in response.headers["content-disposition"]
    zip_path = tmp_path / "empty-attachments.zip"
    zip_path.write_bytes(response.content)
    with ZipFile(zip_path) as attachments_zip:
        assert attachments_zip.namelist() == []


def test_project_attachment_file_routes_do_not_return_files_outside_data_dir(tmp_path):
    client = make_client(tmp_path)
    login(client)
    outside_file = tmp_path / "outside-project.pdf"
    outside_file.write_bytes(b"%PDF-1.7 outside project")
    client.post(
        "/projects",
        data={"year": "2027", "name": "越界读取项目", "budget_amount": "50000"},
    )
    with Session(client.app.state.engine) as session:
        attachment = Attachment(
            project_id=1,
            kind=AttachmentKind.PROCUREMENT_BASIS,
            original_filename="外部项目文件.pdf",
            stored_path="../outside-project.pdf",
        )
        session.add(attachment)
        session.commit()
        session.refresh(attachment)
        attachment_id = attachment.id

    preview_response = client.get(f"/attachments/{attachment_id}/preview")
    inline_response = client.get(f"/attachments/{attachment_id}/file")
    download_response = client.get(f"/attachments/{attachment_id}/download")

    assert preview_response.status_code == 200
    assert preview_response.content != outside_file.read_bytes()
    assert inline_response.content != outside_file.read_bytes()
    assert download_response.content != outside_file.read_bytes()


def test_annual_attachment_file_routes_do_not_return_files_outside_data_dir(tmp_path):
    client = make_client(tmp_path)
    login(client)
    outside_file = tmp_path / "outside-annual.pdf"
    outside_file.write_bytes(b"%PDF-1.7 outside annual")
    client.post(
        "/projects",
        data={"year": "2027", "name": "年度越界读取项目", "budget_amount": "50000"},
    )
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()
        attachment = AnnualAttachment(
            execution_id=execution.id,
            kind=AttachmentKind.ACCEPTANCE,
            original_filename="外部年度文件.pdf",
            stored_path="../outside-annual.pdf",
        )
        session.add(attachment)
        session.commit()
        session.refresh(attachment)
        attachment_id = attachment.id

    preview_response = client.get(f"/annual-attachments/{attachment_id}/preview")
    inline_response = client.get(f"/annual-attachments/{attachment_id}/file")
    download_response = client.get(f"/annual-attachments/{attachment_id}/download")

    assert preview_response.status_code == 200
    assert preview_response.content != outside_file.read_bytes()
    assert inline_response.content != outside_file.read_bytes()
    assert download_response.content != outside_file.read_bytes()


def test_project_attachment_can_be_deleted_from_detail_and_preview(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "附件误传项目", "budget_amount": "50000"},
    )
    client.post(
        "/projects/1/attachments",
        data={"kind": "procurement_basis"},
        files={"file": ("错误采购依据.pdf", b"%PDF-1.7 fake basis", "application/pdf")},
    )
    with Session(client.app.state.engine) as session:
        attachment = session.exec(select(Attachment)).one()
        stored_path = tmp_path / "data" / attachment.stored_path

    detail_response = client.get("/projects/1?return_year=2027&q=误传&status=incomplete")
    preview_response = client.get(
        "/attachments/1/preview?return_year=2027&q=误传&status=incomplete"
    )
    delete_response = client.post(
        "/attachments/1/delete",
        data={"return_year": "2027", "q": "误传", "status": "incomplete"},
        follow_redirects=False,
    )
    with Session(client.app.state.engine) as session:
        deleted_attachment = session.get(Attachment, 1)

    assert detail_response.status_code == 200
    assert 'action="/attachments/1/delete"' in detail_response.text
    assert "return confirm(" in detail_response.text
    assert preview_response.status_code == 200
    assert 'action="/attachments/1/delete"' in preview_response.text
    assert delete_response.status_code == 303
    assert (
        delete_response.headers["location"]
        == "/projects/1?return_year=2027&q=%E8%AF%AF%E4%BC%A0&status=incomplete"
    )
    assert deleted_attachment is None
    assert not stored_path.exists()


def test_ledger_attachment_preview_links_preserve_filter_context(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "附件筛选项目", "budget_amount": "50000"},
    )
    client.post(
        "/projects/1/attachments",
        data={"kind": "procurement_basis"},
        files={"file": ("采购依据.pdf", b"%PDF-1.7 fake basis", "application/pdf")},
    )

    ledger_response = client.get("/projects?year=2027&q=附件&status=incomplete")
    preview_response = client.get(
        "/attachments/1/preview?return_year=2027&q=附件&status=incomplete"
    )

    assert (
        'href="/attachments/1/preview?return_year=2027&amp;q=%E9%99%84%E4%BB%B6&amp;status=incomplete"'
        in ledger_response.text
    )
    assert (
        'href="/projects/1?return_year=2027&amp;q=%E9%99%84%E4%BB%B6&amp;status=incomplete"'
        in preview_response.text
    )


def test_project_list_can_filter_by_year(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "2027 年项目", "budget_amount": "100"},
    )
    client.post(
        "/projects",
        data={"year": "2028", "name": "2028 年项目", "budget_amount": "200"},
    )

    all_response = client.get("/projects")
    filtered_response = client.get("/projects?year=2027")

    assert all_response.status_code == 200
    assert 'name="year"' in all_response.text
    assert '<option value="">全部年度</option>' in all_response.text
    assert '<option value="2027"' in all_response.text
    assert '<option value="2028"' in all_response.text
    assert '<button class="ghost" type="submit">查询</button>' in all_response.text
    assert '<button class="ghost" type="submit">筛选</button>' not in all_response.text
    assert "2027 年项目" in filtered_response.text
    assert "2028 年项目" not in filtered_response.text
    assert '<option value="2027" selected>2027 年</option>' in filtered_response.text
    assert '<select name="year" onchange="this.form.submit()">' in filtered_response.text


def test_project_list_invalid_year_filter_should_fall_back_to_all_projects(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "异常筛选保护项目", "budget_amount": "100"},
    )

    response = client.get("/projects?year=bad")

    assert response.status_code == 200
    assert "异常筛选保护项目" in response.text


def test_project_ledger_money_and_contract_period_are_polished(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "金额样式项目", "budget_amount": "12345"},
    )
    client.post(
        "/projects/1/edit",
        data={
            "year": "2027",
            "name": "金额样式项目",
            "budget_amount": "12345",
            "contract_amount": "12345",
            "contract_start": "2026-05-22",
            "contract_end": "2027-05-21",
            "notes": "",
        },
    )

    response = client.get("/projects?year=2027")
    styles = Path("app/static/styles.css").read_text(encoding="utf-8")

    assert '<span class="amount-value">6,172.50</span>' in response.text
    assert '<span class="amount-badge">' not in response.text
    assert '<span class="period-date">2026-05-22 <span class="period-separator">至</span></span>' in response.text
    assert '<span class="period-separator">至</span>' in response.text
    assert '<span class="period-date">2027-05-21</span>' in response.text
    assert ".amount-value" in styles
    assert ".amount-badge" not in styles
    assert ".period-separator" in styles


def test_project_list_treats_empty_year_filter_as_all_years(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "2027 年项目", "budget_amount": "100"},
    )
    client.post(
        "/projects",
        data={"year": "2028", "name": "2028 年项目", "budget_amount": "200"},
    )

    response = client.get("/projects?year=")

    assert response.status_code == 200
    assert "全部年度" in response.text
    assert "2027 年项目" in response.text
    assert "2028 年项目" in response.text


def test_project_list_search_status_filter_summary_and_money_format(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={
            "year": "2027",
            "name": "网络安全维保",
            "budget_amount": "12345",
            "notes": "堡垒机年度服务",
        },
    )
    client.post(
        "/projects",
        data={"year": "2027", "name": "付款完成项目", "budget_amount": "1000"},
    )
    client.post(
        "/projects/2/edit",
        data={
            "year": "2027",
            "name": "付款完成项目",
            "budget_amount": "1000",
            "contract_amount": "900",
            "contract_start": "2027-01-01",
            "contract_end": "2027-12-31",
            "acceptance_date": "2027-12-31",
            "payment_date": "2028-01-15",
            "notes": "",
        },
    )
    client.post(
        "/projects/2/attachments",
        data={"kind": "signed_contract"},
        files={"file": ("盖章合同.pdf", b"%PDF-1.7 fake", "application/pdf")},
    )
    with Session(client.app.state.engine) as session:
        completed_execution = session.exec(
            select(AnnualExecution).where(AnnualExecution.project_id == 2)
        ).one()
    for kind, filename in [
        ("acceptance", "验收单.pdf"),
        ("invoice", "发票.pdf"),
    ]:
        client.post(
            f"/annual-executions/{completed_execution.id}/attachments",
            data={"kind": kind},
            files={"file": (filename, b"%PDF-1.7 fake", "application/pdf")},
        )
    client.post(
        "/projects",
        data={"year": "2027", "name": "未签合同项目", "budget_amount": "500"},
    )
    client.post(
        "/projects",
        data={"year": "2027", "name": "待验收项目", "budget_amount": "600"},
    )
    client.post(
        "/projects/4/edit",
        data={
            "year": "2027",
            "name": "待验收项目",
            "budget_amount": "600",
            "contract_amount": "580",
            "contract_start": "2027-02-01",
            "contract_end": "2027-12-31",
            "acceptance_date": "",
            "payment_date": "",
            "notes": "",
        },
    )
    client.post(
        "/projects/4/attachments",
        data={"kind": "signed_contract"},
        files={"file": ("待验收盖章合同.pdf", b"%PDF-1.7 fake", "application/pdf")},
    )
    client.post(
        "/projects",
        data={"year": "2027", "name": "待付款项目", "budget_amount": "700"},
    )
    client.post(
        "/projects/5/edit",
        data={
            "year": "2027",
            "name": "待付款项目",
            "budget_amount": "700",
            "contract_amount": "680",
            "contract_start": "2027-03-01",
            "contract_end": "2027-12-31",
            "acceptance_date": "2027-12-31",
            "payment_date": "",
            "notes": "",
        },
    )
    client.post(
        "/projects/5/attachments",
        data={"kind": "signed_contract"},
        files={"file": ("待付款盖章合同.pdf", b"%PDF-1.7 fake", "application/pdf")},
    )
    with Session(client.app.state.engine) as session:
        unpaid_execution = session.exec(
            select(AnnualExecution).where(AnnualExecution.project_id == 5)
        ).one()
    client.post(
        f"/annual-executions/{unpaid_execution.id}/attachments",
        data={"kind": "acceptance"},
        files={"file": ("待付款验收单.pdf", b"%PDF-1.7 fake", "application/pdf")},
    )

    response = client.get("/projects?q=堡垒机&status=incomplete")
    completed_response = client.get("/projects?status=completed")
    unsigned_response = client.get("/projects?q=未签&status=unsigned_contract")
    unaccepted_response = client.get("/projects?q=待验收&status=unaccepted")
    unpaid_response = client.get("/projects?q=待付款&status=unpaid")

    assert response.status_code == 200
    assert 'name="q"' in response.text
    assert 'value="堡垒机"' in response.text
    assert '<option value="incomplete" selected>资料不完整</option>' in response.text
    assert 'href="/projects/export?q=' in response.text
    assert "status=incomplete" in response.text
    assert "网络安全维保" in response.text
    assert "付款完成项目" not in response.text
    assert "当前项目" in response.text
    assert "预算合计" in response.text
    assert "12,345.00" in response.text
    assert "待补资料" in response.text
    assert "付款完成项目" in completed_response.text
    assert "网络安全维保" not in completed_response.text
    assert "未签合同项目" in unsigned_response.text
    assert "待验收项目" in unaccepted_response.text
    assert "待付款项目" in unpaid_response.text


def test_project_list_searches_annual_execution_notes(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={
            "year": "2027",
            "name": "年度备注检索项目",
            "budget_amount": "12000",
            "notes": "项目主档没有目标词",
        },
    )
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()
    client.post(
        f"/annual-executions/{execution.id}/edit",
        data={
            "budget_amount": "12000",
            "contract_amount": "",
            "acceptance_date": "",
            "payment_date": "",
            "notes": "年度执行备注包含堡垒机续保",
        },
    )

    response = client.get("/projects?q=堡垒机续保")

    assert response.status_code == 200
    assert "年度备注检索项目" in response.text
    assert 'value="堡垒机续保"' in response.text


def test_project_list_exports_filtered_year_to_excel(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "2027 年项目", "budget_amount": "100"},
    )
    client.post(
        "/projects",
        data={"year": "2028", "name": "2028 年项目", "budget_amount": "200"},
    )

    page_response = client.get("/projects?year=2027")
    export_response = client.get("/projects/export?year=2027")

    assert page_response.status_code == 200
    assert 'href="/projects/export?year=2027"' in page_response.text
    assert export_response.status_code == 200
    assert export_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "projects-2027.xlsx" in export_response.headers["content-disposition"]

    workbook_path = tmp_path / "export.xlsx"
    workbook_path.write_bytes(export_response.content)
    with ZipFile(workbook_path) as workbook_zip:
        workbook_xml = "\n".join(
            workbook_zip.read(name).decode("utf-8", errors="ignore")
            for name in workbook_zip.namelist()
            if name.startswith("xl/")
            and name.endswith(".xml")
            and ("worksheet" in name or "sharedStrings" in name)
        )
    assert "2027 年项目" in workbook_xml
    assert "2028 年项目" not in workbook_xml


def test_project_excel_export_uses_search_and_status_filters(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={
            "year": "2027",
            "name": "堡垒机维保",
            "budget_amount": "12345",
            "notes": "网络安全",
        },
    )
    client.post(
        "/projects",
        data={"year": "2027", "name": "机房巡检", "budget_amount": "200"},
    )

    export_response = client.get("/projects/export?year=2027&q=网络&status=incomplete")

    assert export_response.status_code == 200
    workbook_path = tmp_path / "export.xlsx"
    workbook_path.write_bytes(export_response.content)
    with ZipFile(workbook_path) as workbook_zip:
        workbook_xml = "\n".join(
            workbook_zip.read(name).decode("utf-8", errors="ignore")
            for name in workbook_zip.namelist()
            if name.startswith("xl/")
            and name.endswith(".xml")
            and ("worksheet" in name or "sharedStrings" in name)
        )
    assert "堡垒机维保" in workbook_xml
    assert "机房巡检" not in workbook_xml
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    assert sheet["C3"].value == 12345
    assert sheet["C3"].number_format == "#,##0.00"
    assert sheet["D3"].number_format == "#,##0.00"


def test_project_excel_export_omits_attachment_columns(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "2027 年项目", "budget_amount": "100"},
    )
    client.post(
        "/projects/1/attachments",
        data={"kind": "procurement_basis"},
        files={"file": ("采购依据.pdf", b"%PDF-1.7 fake basis", "application/pdf")},
    )

    export_response = client.get("/projects/export?year=2027")

    workbook_path = tmp_path / "export.xlsx"
    workbook_path.write_bytes(export_response.content)
    with ZipFile(workbook_path) as workbook_zip:
        workbook_xml = "\n".join(
            workbook_zip.read(name).decode("utf-8", errors="ignore")
            for name in workbook_zip.namelist()
            if name.startswith("xl/")
            and name.endswith(".xml")
            and ("worksheet" in name or "sharedStrings" in name)
        )
    assert "采购依据.pdf" not in workbook_xml
    assert "合同审签 PDF" not in workbook_xml
    assert "其他附件" not in workbook_xml


def test_cross_year_project_appears_in_each_execution_year(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={
            "year": "2026",
            "name": "三年云平台服务",
            "budget_amount": "300000",
        },
    )
    client.post(
        "/projects/1/edit",
        data={
            "year": "2026",
            "name": "三年云平台服务",
            "budget_amount": "300000",
            "contract_amount": "270000",
            "contract_start": "2027-01-01",
            "contract_end": "2029-12-31",
            "notes": "",
        },
    )

    contract_year_response = client.get("/projects?year=2026")
    execution_year_response = client.get("/projects?year=2028")

    assert contract_year_response.status_code == 200
    assert "三年云平台服务" in contract_year_response.text
    assert "合同管理年" in contract_year_response.text
    assert "100,000.00" not in contract_year_response.text
    assert execution_year_response.status_code == 200
    assert "三年云平台服务" in execution_year_response.text
    assert "100,000.00" in execution_year_response.text
    assert "90,000.00" in execution_year_response.text


def test_project_detail_back_link_preserves_ledger_year_filter(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "次年管理项目", "budget_amount": "30000"},
    )
    client.post(
        "/projects/1/edit",
        data={
            "year": "2027",
            "name": "次年管理项目",
            "budget_amount": "30000",
            "contract_amount": "30000",
            "contract_start": "2026-05-22",
            "contract_end": "2027-05-22",
            "notes": "",
        },
    )

    ledger_response = client.get("/projects?year=2026")
    detail_response = client.get("/projects/1?return_year=2026")

    assert ledger_response.status_code == 200
    assert 'href="/projects/1?return_year=2026"' in ledger_response.text
    assert detail_response.status_code == 200
    assert 'href="/projects?year=2026"' in detail_response.text


def test_project_edit_rejects_contract_start_after_contract_end(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "合同日期保护项目", "budget_amount": "30000"},
    )

    response = client.post(
        "/projects/1/edit",
        data={
            "year": "2027",
            "name": "合同日期保护项目",
            "budget_amount": "30000",
            "contract_amount": "30000",
            "contract_start": "2028-01-01",
            "contract_end": "2027-12-31",
            "notes": "",
        },
        follow_redirects=False,
    )
    detail_response = client.get("/projects/1")
    with Session(client.app.state.engine) as session:
        project = session.get(Project, 1)

    assert response.status_code == 400
    assert "合同开始日期不能晚于合同结束日期" in response.text
    assert project.contract_start is None
    assert project.contract_end is None
    assert "2028-01-01" not in detail_response.text


def test_project_edit_rejects_blank_name_and_invalid_date(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "编辑保护项目", "budget_amount": "30000"},
    )

    blank_name_response = client.post(
        "/projects/1/edit",
        data={
            "year": "2027",
            "name": "   ",
            "budget_amount": "30000",
            "contract_amount": "",
            "contract_start": "",
            "contract_end": "",
            "notes": "",
        },
        follow_redirects=False,
    )
    invalid_date_response = client.post(
        "/projects/1/edit",
        data={
            "year": "2027",
            "name": "编辑保护项目",
            "budget_amount": "30000",
            "contract_amount": "",
            "contract_start": "bad-date",
            "contract_end": "",
            "notes": "",
        },
        follow_redirects=False,
    )
    with Session(client.app.state.engine) as session:
        project = session.get(Project, 1)

    assert blank_name_response.status_code == 400
    assert "项目名称不能为空" in blank_name_response.text
    assert invalid_date_response.status_code == 400
    assert "合同开始日期格式不正确" in invalid_date_response.text
    assert project.name == "编辑保护项目"
    assert project.contract_start is None


def test_project_edit_rejects_non_finite_contract_amount(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "合同金额保护项目", "budget_amount": "30000"},
    )

    response = client.post(
        "/projects/1/edit",
        data={
            "year": "2027",
            "name": "合同金额保护项目",
            "budget_amount": "30000",
            "contract_amount": "inf",
            "contract_start": "",
            "contract_end": "",
            "notes": "",
        },
        follow_redirects=False,
    )
    with Session(client.app.state.engine) as session:
        project = session.get(Project, 1)

    assert response.status_code == 400
    assert "合同总额格式不正确" in response.text
    assert project.contract_amount is None


def test_project_detail_save_forms_preserve_scroll_position(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "滚动保持项目", "budget_amount": "30000"},
    )

    base_response = client.get("/projects/1")
    detail_response = client.get("/projects/1")

    assert 'src="/static/app.js"' in base_response.text
    assert detail_response.text.count("data-preserve-scroll") == 5
    assert 'action="/projects/1/edit" class="form-grid" data-preserve-scroll' in detail_response.text
    assert (
        'action="/projects/1/attachments" enctype="multipart/form-data" '
        'class="upload-row" data-preserve-scroll'
    ) in detail_response.text
    assert "/annual-executions/1/edit" in detail_response.text
    assert "/annual-executions/1/attachments" in detail_response.text
    assert 'action="/projects/1/delete"' in detail_response.text
    delete_form_start = detail_response.text.index('action="/projects/1/delete"')
    delete_form_end = detail_response.text.index("</form>", delete_form_start)
    assert "data-preserve-scroll" not in detail_response.text[
        delete_form_start:delete_form_end
    ]


def test_project_detail_delete_preserves_full_ledger_filter_context(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "筛选删除项目", "budget_amount": "30000"},
    )
    client.post(
        "/projects",
        data={"year": "2027", "name": "筛选保留项目", "budget_amount": "20000"},
    )

    ledger_response = client.get("/projects?year=2027&q=删除&status=incomplete")
    detail_response = client.get(
        "/projects/1?return_year=2027&q=删除&status=incomplete"
    )
    delete_response = client.post(
        "/projects/1/delete",
        data={"return_year": "2027", "q": "删除", "status": "incomplete"},
        follow_redirects=False,
    )

    assert ledger_response.status_code == 200
    assert (
        'href="/projects/1?return_year=2027&amp;q=%E5%88%A0%E9%99%A4&amp;status=incomplete"'
        in ledger_response.text
    )
    assert detail_response.status_code == 200
    assert (
        'href="/projects?year=2027&amp;q=%E5%88%A0%E9%99%A4&amp;status=incomplete"'
        in detail_response.text
    )
    assert 'name="q" value="删除"' in detail_response.text
    assert 'name="status" value="incomplete"' in detail_response.text
    assert (
        delete_response.headers["location"]
        == "/projects?year=2027&q=%E5%88%A0%E9%99%A4&status=incomplete"
    )


def test_project_detail_save_redirect_preserves_return_year(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "返回年度项目", "budget_amount": "30000"},
    )

    response = client.post(
        "/projects/1/edit",
        data={
            "year": "2027",
            "name": "返回年度项目",
            "budget_amount": "30000",
            "contract_amount": "30000",
            "contract_start": "2026-05-22",
            "contract_end": "2027-05-22",
            "notes": "",
            "return_year": "2026",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"] == "/projects/1?return_year=2026"


def test_annual_execution_save_redirect_preserves_return_year(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "年度返回项目", "budget_amount": "30000"},
    )
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()

    response = client.post(
        f"/annual-executions/{execution.id}/edit",
        data={
            "budget_amount": "30000",
            "contract_amount": "30000",
            "acceptance_date": "",
            "payment_date": "",
            "notes": "",
            "return_year": "2026",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"] == "/projects/1?return_year=2026"


def test_annual_execution_detail_edit_and_annual_attachment_upload(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "年度安全服务", "budget_amount": "60000"},
    )
    detail_response = client.get("/projects/1")
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()

    assert detail_response.status_code == 200
    assert "年度执行计划" in detail_response.text
    assert 'name="contract_only"' in detail_response.text

    client.post(
        f"/annual-executions/{execution.id}/edit",
        data={
            "budget_amount": "62000",
            "contract_amount": "61000",
            "acceptance_date": "2027-12-31",
            "payment_date": "2028-01-15",
            "notes": "第一年付款",
        },
    )
    client.post(
        f"/annual-executions/{execution.id}/attachments",
        data={"kind": "acceptance"},
        files={"file": ("2027验收单.pdf", b"%PDF-1.7 fake", "application/pdf")},
    )
    client.post(
        f"/annual-executions/{execution.id}/attachments",
        data={"kind": "invoice"},
        files={"file": ("2027发票.pdf", b"%PDF-1.7 fake", "application/pdf")},
    )
    with Session(client.app.state.engine) as session:
        annual_attachments = session.exec(select(AnnualAttachment)).all()
    preview_response = client.get(f"/annual-attachments/{annual_attachments[0].id}/preview")
    inline_response = client.get(f"/annual-attachments/{annual_attachments[0].id}/file")
    download_response = client.get(f"/annual-attachments/{annual_attachments[0].id}/download")
    year_zip_response = client.get(
        f"/annual-executions/{execution.id}/attachments/download-year"
    )
    all_zip_response = client.get("/projects/1/attachments/download-all")

    updated_response = client.get("/projects/1")

    assert preview_response.status_code == 200
    assert f'src="/annual-attachments/{annual_attachments[0].id}/file"' in preview_response.text
    assert inline_response.headers["content-type"] == "application/pdf"
    assert download_response.headers["content-type"] == "application/pdf"
    assert "project-1-2027-attachments.zip" in year_zip_response.headers["content-disposition"]
    assert "project-1-attachments.zip" in all_zip_response.headers["content-disposition"]
    assert "62,000.00" in updated_response.text
    assert "2027验收单.pdf" in updated_response.text
    assert "2027发票.pdf" in updated_response.text


def test_annual_execution_edit_rejects_invalid_amount_and_date(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "年度输入保护项目", "budget_amount": "60000"},
    )
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()

    invalid_amount_response = client.post(
        f"/annual-executions/{execution.id}/edit",
        data={
            "budget_amount": "abc",
            "contract_amount": "",
            "acceptance_date": "",
            "payment_date": "",
            "notes": "",
        },
        follow_redirects=False,
    )
    invalid_date_response = client.post(
        f"/annual-executions/{execution.id}/edit",
        data={
            "budget_amount": "60000",
            "contract_amount": "",
            "acceptance_date": "",
            "payment_date": "bad-date",
            "notes": "",
        },
        follow_redirects=False,
    )
    with Session(client.app.state.engine) as session:
        updated_execution = session.get(AnnualExecution, execution.id)

    assert invalid_amount_response.status_code == 400
    assert "年度预算格式不正确" in invalid_amount_response.text
    assert invalid_date_response.status_code == 400
    assert "付款日期格式不正确" in invalid_date_response.text
    assert updated_execution.budget_amount == 60000
    assert updated_execution.payment_date is None


def test_annual_execution_edit_rejects_non_finite_budget_amount(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "年度非有限金额项目", "budget_amount": "60000"},
    )
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()

    response = client.post(
        f"/annual-executions/{execution.id}/edit",
        data={
            "budget_amount": "inf",
            "contract_amount": "",
            "acceptance_date": "",
            "payment_date": "",
            "notes": "",
        },
        follow_redirects=False,
    )
    with Session(client.app.state.engine) as session:
        updated_execution = session.get(AnnualExecution, execution.id)

    assert response.status_code == 400
    assert "年度预算格式不正确" in response.text
    assert updated_execution.budget_amount == 60000


def test_project_attachment_upload_rejects_annual_attachment_kind(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "错误附件类型项目", "budget_amount": "60000"},
    )

    response = client.post(
        "/projects/1/attachments",
        data={"kind": "invoice"},
        files={"file": ("发票.pdf", b"%PDF-1.7 fake", "application/pdf")},
        follow_redirects=False,
    )

    assert response.status_code == 400
    assert "项目级附件只支持" in response.text


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("采购依据.pdf", b"not a pdf"),
        ("采购依据.txt", b"%PDF-1.7 fake"),
    ],
)
def test_project_attachment_upload_invalid_pdf_returns_controlled_error(
    tmp_path,
    filename,
    content,
):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "非法项目级附件项目", "budget_amount": "60000"},
    )

    response = client.post(
        "/projects/1/attachments",
        data={"kind": "procurement_basis"},
        files={"file": (filename, content, "application/pdf")},
        follow_redirects=False,
    )

    assert response.status_code == 400
    assert "只支持上传 PDF 文件" in response.text


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("验收单.pdf", b"not a pdf"),
        ("验收单.txt", b"%PDF-1.7 fake"),
    ],
)
def test_annual_attachment_upload_invalid_pdf_returns_controlled_error(
    tmp_path,
    filename,
    content,
):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "非法年度附件项目", "budget_amount": "60000"},
    )
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()

    response = client.post(
        f"/annual-executions/{execution.id}/attachments",
        data={"kind": "acceptance"},
        files={"file": (filename, content, "application/pdf")},
        follow_redirects=False,
    )

    assert response.status_code == 400
    assert "只支持上传 PDF 文件" in response.text


def test_full_business_flow_reaches_completed_status_and_exports_artifacts(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={
            "year": "2027",
            "name": "完整闭环项目",
            "budget_amount": "50000",
            "notes": "年度预算项目",
        },
    )
    client.post(
        "/projects/1/edit",
        data={
            "year": "2027",
            "name": "完整闭环项目",
            "budget_amount": "50000",
            "contract_amount": "48000",
            "contract_start": "2027-01-01",
            "contract_end": "2027-12-31",
            "notes": "年度预算项目",
        },
    )
    for kind, filename in [
        ("procurement_basis", "采购依据.pdf"),
        ("contract_review", "合同审签.pdf"),
        ("signed_contract", "盖章合同.pdf"),
        ("other", "补充说明.pdf"),
    ]:
        client.post(
            "/projects/1/attachments",
            data={"kind": kind},
            files={"file": (filename, b"%PDF-1.7 fake", "application/pdf")},
        )
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()

    client.post(
        f"/annual-executions/{execution.id}/edit",
        data={
            "budget_amount": "50000",
            "contract_amount": "48000",
            "acceptance_date": "2027-12-31",
            "payment_date": "2028-01-15",
            "notes": "按合同完成验收付款",
        },
    )
    for kind, filename in [
        ("acceptance", "验收单.pdf"),
        ("invoice", "发票.pdf"),
    ]:
        client.post(
            f"/annual-executions/{execution.id}/attachments",
            data={"kind": kind},
            files={"file": (filename, b"%PDF-1.7 fake", "application/pdf")},
        )

    completed_response = client.get("/projects?year=2027&status=completed")
    excel_response = client.get("/projects/export?year=2027&status=completed")
    all_zip_response = client.get("/projects/1/attachments/download-all")
    year_zip_response = client.get(
        f"/annual-executions/{execution.id}/attachments/download-year"
    )

    assert completed_response.status_code == 200
    assert "完整闭环项目" in completed_response.text
    assert "合同已签" in completed_response.text
    assert "已验收" in completed_response.text
    assert "已付款" in completed_response.text

    workbook_path = tmp_path / "completed-export.xlsx"
    workbook_path.write_bytes(excel_response.content)
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    assert sheet["A3"].value == 2027
    assert sheet["B3"].value == "完整闭环项目"
    assert sheet["C3"].value == 50000
    assert sheet["D3"].value == 48000
    assert sheet["F3"].number_format == "yyyy-mm-dd"
    assert sheet["G3"].number_format == "yyyy-mm-dd"
    assert sheet["H3"].number_format == "yyyy-mm-dd"
    assert sheet["I3"].number_format == "yyyy-mm-dd"

    all_zip_path = tmp_path / "all-attachments.zip"
    all_zip_path.write_bytes(all_zip_response.content)
    with ZipFile(all_zip_path) as attachments_zip:
        names = attachments_zip.namelist()
        assert any(name.endswith("采购依据.pdf") for name in names)
        assert any(name.endswith("合同审签.pdf") for name in names)
        assert any(name.endswith("盖章合同.pdf") for name in names)
        assert any(name.endswith("验收单.pdf") for name in names)
        assert any(name.endswith("发票.pdf") for name in names)

    year_zip_path = tmp_path / "year-attachments.zip"
    year_zip_path.write_bytes(year_zip_response.content)
    with ZipFile(year_zip_path) as attachments_zip:
        names = attachments_zip.namelist()
        assert any(name.startswith("项目资料/采购依据/") for name in names)
        assert any(name.startswith("年度资料/验收单/") for name in names)
        assert any(name.startswith("年度资料/发票/") for name in names)


def test_annual_attachment_upload_controls_use_aligned_grid(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "年度上传对齐项目", "budget_amount": "60000"},
    )

    response = client.get("/projects/1")
    styles = Path("app/static/styles.css").read_text(encoding="utf-8")

    assert 'class="annual-upload-kind"' in response.text
    assert 'id="annual-file-1-acceptance"' in response.text
    assert 'for="annual-file-1-acceptance"' in response.text
    assert ".annual-uploads .compact-upload" in styles
    assert "grid-template-columns: auto minmax(160px, 1fr) auto;" in styles
    assert "grid-auto-flow: column;" in styles


def test_annual_attachment_can_be_deleted_from_detail_and_preview(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "年度附件误传项目", "budget_amount": "60000"},
    )
    with Session(client.app.state.engine) as session:
        execution = session.exec(select(AnnualExecution)).one()
    client.post(
        f"/annual-executions/{execution.id}/attachments",
        data={"kind": "acceptance"},
        files={"file": ("错误验收单.pdf", b"%PDF-1.7 fake", "application/pdf")},
    )
    with Session(client.app.state.engine) as session:
        attachment = session.exec(select(AnnualAttachment)).one()
        stored_path = tmp_path / "data" / attachment.stored_path

    detail_response = client.get("/projects/1?return_year=2027&q=误传&status=incomplete")
    preview_response = client.get(
        "/annual-attachments/1/preview?return_year=2027&q=误传&status=incomplete"
    )
    delete_response = client.post(
        "/annual-attachments/1/delete",
        data={"return_year": "2027", "q": "误传", "status": "incomplete"},
        follow_redirects=False,
    )
    with Session(client.app.state.engine) as session:
        deleted_attachment = session.get(AnnualAttachment, 1)

    assert detail_response.status_code == 200
    assert 'action="/annual-attachments/1/delete"' in detail_response.text
    assert "return confirm(" in detail_response.text
    assert preview_response.status_code == 200
    assert 'action="/annual-attachments/1/delete"' in preview_response.text
    assert delete_response.status_code == 303
    assert (
        delete_response.headers["location"]
        == "/projects/1?return_year=2027&q=%E8%AF%AF%E4%BC%A0&status=incomplete"
    )
    assert deleted_attachment is None
    assert not stored_path.exists()


def test_project_detail_can_delete_project_after_confirmation(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "待删除项目", "budget_amount": "30000"},
    )

    detail_response = client.get("/projects/1?return_year=2027")
    delete_response = client.post(
        "/projects/1/delete",
        data={"return_year": "2027"},
        follow_redirects=False,
    )
    with Session(client.app.state.engine) as session:
        deleted_project = session.get(Project, 1)

    assert detail_response.status_code == 200
    assert "onsubmit=" in detail_response.text
    assert "return confirm(" in detail_response.text
    assert "删除项目" in detail_response.text
    assert delete_response.status_code == 303
    assert delete_response.headers["location"] == "/projects?year=2027"
    assert deleted_project is None


def test_deleting_missing_project_or_attachment_redirects_safely(tmp_path):
    client = make_client(tmp_path)
    login(client)

    missing_project_response = client.post(
        "/projects/999/delete",
        follow_redirects=False,
    )
    missing_attachment_response = client.post(
        "/attachments/999/delete",
        follow_redirects=False,
    )
    missing_annual_attachment_response = client.post(
        "/annual-attachments/999/delete",
        follow_redirects=False,
    )
    missing_preview_response = client.get(
        "/attachments/999/preview",
        follow_redirects=False,
    )

    assert missing_project_response.status_code == 303
    assert missing_project_response.headers["location"] == "/projects"
    assert missing_attachment_response.status_code == 303
    assert missing_attachment_response.headers["location"] == "/projects"
    assert missing_annual_attachment_response.status_code == 303
    assert missing_annual_attachment_response.headers["location"] == "/projects"
    assert missing_preview_response.status_code == 303
    assert missing_preview_response.headers["location"] == "/projects"


def test_project_ledger_can_delete_project_and_keep_filters(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "台账删除项目", "budget_amount": "30000"},
    )
    client.post(
        "/projects",
        data={"year": "2027", "name": "保留项目", "budget_amount": "20000"},
    )

    ledger_response = client.get("/projects?year=2027&q=删除&status=incomplete")
    delete_response = client.post(
        "/projects/1/delete",
        data={"return_year": "2027", "q": "删除", "status": "incomplete"},
        follow_redirects=False,
    )
    with Session(client.app.state.engine) as session:
        deleted_project = session.get(Project, 1)
        remaining_projects = session.exec(select(Project)).all()

    assert ledger_response.status_code == 200
    assert "<th>操作</th>" in ledger_response.text
    assert 'action="/projects/1/delete"' in ledger_response.text
    assert "return confirm(" in ledger_response.text
    assert 'name="return_year" value="2027"' in ledger_response.text
    assert 'name="q" value="删除"' in ledger_response.text
    assert 'name="status" value="incomplete"' in ledger_response.text
    assert "删除" in ledger_response.text
    assert delete_response.status_code == 303
    assert (
        delete_response.headers["location"]
        == "/projects?year=2027&q=%E5%88%A0%E9%99%A4&status=incomplete"
    )
    assert deleted_project is None
    assert [project.name for project in remaining_projects] == ["保留项目"]


def test_manual_contract_management_checkbox_updates_ledger(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2027", "name": "次年付款合同", "budget_amount": "30000"},
    )
    client.post(
        "/projects/1/edit",
        data={
            "year": "2027",
            "name": "次年付款合同",
            "budget_amount": "30000",
            "contract_amount": "30000",
            "contract_start": "2026-05-22",
            "contract_end": "2027-05-22",
            "notes": "",
        },
    )
    with Session(client.app.state.engine) as session:
        execution = session.exec(
            select(AnnualExecution).where(AnnualExecution.year == 2026)
        ).one()

    client.post(
        f"/annual-executions/{execution.id}/edit",
        data={
            "budget_amount": "",
            "contract_amount": "15000",
            "contract_only": "on",
            "acceptance_date": "",
            "payment_date": "",
            "notes": "本年只管理合同",
        },
    )

    ledger_response = client.get("/projects?year=2026")
    next_year_response = client.get("/projects?year=2027")

    assert ledger_response.status_code == 200
    assert "次年付款合同" in ledger_response.text
    assert "合同管理年" in ledger_response.text
    assert "15,000.00" not in ledger_response.text
    assert "30,000.00" in next_year_response.text


def test_excel_export_uses_annual_execution_rows(tmp_path):
    client = make_client(tmp_path)
    login(client)
    client.post(
        "/projects",
        data={"year": "2026", "name": "三年云平台服务", "budget_amount": "300000"},
    )
    client.post(
        "/projects/1/edit",
        data={
            "year": "2026",
            "name": "三年云平台服务",
            "budget_amount": "300000",
            "contract_amount": "270000",
            "contract_start": "2027-01-01",
            "contract_end": "2029-12-31",
            "notes": "",
        },
    )

    export_response = client.get("/projects/export?year=2028")

    workbook_path = tmp_path / "annual-export.xlsx"
    workbook_path.write_bytes(export_response.content)
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    assert sheet["A3"].value == 2028
    assert sheet["B3"].value == "三年云平台服务"
    assert sheet["C3"].value == 100000
    assert sheet["D3"].value == 90000


def test_runtime_configuration_docs_and_start_script_match_mvp_contract():
    env_example = Path(".env.example").read_text(encoding="utf-8")
    start_script = Path("scripts/start.ps1").read_text(encoding="utf-8")
    readme = Path("README.md").read_text(encoding="utf-8")
    startup_doc = Path("启动停止重启说明.md").read_text(encoding="utf-8")

    for key in [
        "PROJMAN_USERNAME",
        "PROJMAN_PASSWORD",
        "PROJMAN_SECRET_KEY",
        "PROJMAN_DATABASE_URL",
        "PROJMAN_DATA_DIR",
    ]:
        assert key in env_example
        assert key in startup_doc

    assert 'Test-Path ".env"' in start_script
    assert 'uv venv .venv' in start_script
    assert 'uv run uvicorn app.main:app --host 0.0.0.0 --port 8765' in start_script
    assert 'PROJMAN_USERNAME = "admin"' in start_script
    assert 'PROJMAN_PASSWORD = "admin"' in start_script
    assert "http://127.0.0.1:8765" in readme
    assert "data/app.db" in readme
    assert "data/attachments/" in readme
    assert "data/annual_attachments/" in readme
